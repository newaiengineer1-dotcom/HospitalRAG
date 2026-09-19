from pathlib import Path
import json
import hashlib
import re
import csv

import numpy as np
import faiss

from sentence_transformers import SentenceTransformer

from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path.cwd()

DOCUMENT_DIR = BASE_DIR / "hospital_documents"
KB_DIR = BASE_DIR / "knowledge_base"

INDEX_FILE = KB_DIR / "faiss.index"
METADATA_FILE = KB_DIR / "metadata.json"
CONFIG_FILE = KB_DIR / "config.json"
MANIFEST_FILE = KB_DIR / "manifest.json"

EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"

CHUNK_SIZE = 1800
CHUNK_OVERLAP = 300

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".xlsx",
    ".pptx",
}


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(text: str) -> str:
    if not text:
        return ""

    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


# ============================================================
# CHUNKING
# ============================================================

def create_chunks(text: str):
    text = clean_text(text)

    if not text:
        return []

    chunks = []

    start = 0
    text_length = len(text)

    while start < text_length:

         end = min(start + CHUNK_SIZE, text_length)

         # Try to finish at a natural boundary
         if end < text_length:
             boundary = text.rfind("\n\n", start, end)

             if boundary == -1:
                 boundary = text.rfind(". ", start, end)

             if boundary == -1:
                 boundary = text.rfind(" ", start, end)

             if boundary > start + int(CHUNK_SIZE * 0.55):
                 end = boundary + 1

         chunk = text[start:end].strip()

         if chunk:
             chunks.append(chunk)

         if end >= text_length:
             break

         next_start = end - CHUNK_OVERLAP

         if next_start <= start:
             next_start = end

         start = next_start

    return chunks


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_pdf(path: Path):
    records = []

    try:
        reader = PdfReader(str(path))

        for page_number, page in enumerate(reader.pages, start=1):

            text = page.extract_text() or ""

            if text.strip():
                records.append({
                    "text": text,
                    "page_number": page_number,
                    "section": "",
                })

    except Exception as exc:
        print(f"[PDF ERROR] {path.name}: {exc}")

    return records


def extract_docx(path: Path):
    records = []

    try:
        doc = DocxDocument(str(path))

        paragraphs = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                )

                if row_text.strip():
                    paragraphs.append(row_text)

        text = "\n".join(paragraphs)

        if text.strip():
            records.append({
                "text": text,
                "page_number": None,
                "section": "",
            })

    except Exception as exc:
        print(f"[DOCX ERROR] {path.name}: {exc}")

    return records


def extract_text_file(path: Path):
    try:
        text = path.read_text(
            encoding="utf-8",
            errors="ignore"
        )

        if text.strip():
            return [{
                "text": text,
                "page_number": None,
                "section": "",
            }]

    except Exception as exc:
        print(f"[TEXT ERROR] {path.name}: {exc}")

    return []


def extract_csv(path: Path):
    rows = []

    try:
        with open(
            path,
            "r",
            encoding="utf-8",
            errors="ignore",
            newline=""
        ) as file:

            reader = csv.reader(file)

            for row in reader:
                row_text = " | ".join(str(x) for x in row)

                if row_text.strip():
                    rows.append(row_text)

        text = "\n".join(rows)

        if text.strip():
            return [{
                "text": text,
                "page_number": None,
                "section": "",
            }]

    except Exception as exc:
        print(f"[CSV ERROR] {path.name}: {exc}")

    return []


def extract_xlsx(path: Path):
    records = []

    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True
        )

        for worksheet in workbook.worksheets:

            rows = []

            for row in worksheet.iter_rows(values_only=True):

                values = [
                    str(value).strip()
                    for value in row
                    if value is not None
                ]

                if values:
                    rows.append(" | ".join(values))

            text = "\n".join(rows)

            if text.strip():

                records.append({
                    "text": text,
                    "page_number": None,
                    "section": worksheet.title,
                })

    except Exception as exc:
        print(f"[XLSX ERROR] {path.name}: {exc}")

    return records


def extract_pptx(path: Path):
    records = []

    try:
        presentation = Presentation(str(path))

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            texts = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):
                    if shape.text.strip():
                        texts.append(shape.text)

            text = "\n".join(texts)

            if text.strip():
                records.append({
                    "text": text,
                    "page_number": slide_number,
                    "section": f"Slide {slide_number}",
                })

    except Exception as exc:
        print(f"[PPTX ERROR] {path.name}: {exc}")

    return records


def extract_document(path: Path):

    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(path)

    if extension == ".docx":
        return extract_docx(path)

    if extension in {".txt", ".md"}:
        return extract_text_file(path)

    if extension == ".csv":
        return extract_csv(path)

    if extension == ".xlsx":
        return extract_xlsx(path)

    if extension == ".pptx":
        return extract_pptx(path)

    return []


# ============================================================
# DOCUMENT ID
# ============================================================

def make_document_id(filename: str) -> str:

    return hashlib.sha256(
        filename.encode("utf-8")
    ).hexdigest()[:16]


# ============================================================
# BUILD KNOWLEDGE BASE
# ============================================================

def build_knowledge_base():

    DOCUMENT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    KB_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    files = [
        path
        for path in DOCUMENT_DIR.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
    ]

    if not files:
        raise RuntimeError(
            "No supported documents found in "
            f"{DOCUMENT_DIR}"
        )

    print("=" * 60)
    print("HOSPITAL RAG - KNOWLEDGE BASE BUILD")
    print("=" * 60)

    metadata = []
    texts = []

    chunk_counter = 0

    for path in files:

        print(f"\nProcessing: {path.name}")

        document_id = make_document_id(path.name)

        records = extract_document(path)

        for record in records:

            chunks = create_chunks(
                record["text"]
            )

            for chunk_index, chunk_text in enumerate(
                chunks
            ):

                chunk_id = f"chunk_{chunk_counter:06d}"

                metadata.append({
                    "document_id": document_id,
                    "filename": path.name,
                    "page_number": record["page_number"],
                    "section": record["section"],
                    "chunk_id": chunk_id,
                    "chunk_index": chunk_index,
                    "text": chunk_text,
                    "source": path.name,
                })

                texts.append(chunk_text)

                chunk_counter += 1

        print(
            f"  Extracted records: {len(records)}"
        )

    if not texts:
        raise RuntimeError(
            "Documents were found, but no text "
            "could be extracted."
        )

    print("\nTotal chunks:", len(texts))

    # ========================================================
    # EMBEDDINGS
    # ========================================================

    print("\nLoading embedding model:")

    print(EMBEDDING_MODEL)

    model = SentenceTransformer(
        EMBEDDING_MODEL
    )

    print("Creating embeddings...")

    embeddings = model.encode(
        texts,
        batch_size=32,
        show_progress_bar=True,
        normalize_embeddings=True,
        convert_to_numpy=True,
    )

    embeddings = np.asarray(
        embeddings,
        dtype="float32"
    )

    print(
        "Embedding shape:",
        embeddings.shape
    )

    # ========================================================
    # FAISS
    # ========================================================

    dimension = embeddings.shape[1]

    index = faiss.IndexFlatIP(
        dimension
    )

    index.add(embeddings)

    faiss.write_index(
        index,
        str(INDEX_FILE)
    )

    # ========================================================
    # METADATA
    # ========================================================

    with open(
        METADATA_FILE,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            metadata,
            file,
            ensure_ascii=False,
            indent=2
        )

    # ========================================================
    # CONFIG
    # ========================================================

    config = {
        "embedding_model": EMBEDDING_MODEL,
        "vector_dimension": dimension,
        "chunk_size": CHUNK_SIZE,
        "chunk_overlap": CHUNK_OVERLAP,
        "similarity_type": "cosine_via_normalized_inner_product",
        "total_chunks": len(metadata),
    }

    with open(
        CONFIG_FILE,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            config,
            file,
            indent=2
        )

    # ========================================================
    # MANIFEST
    # ========================================================

    manifest = {
        "application": "Hospital Knowledge Assistant",
        "total_documents": len(files),
        "total_chunks": len(metadata),
        "embedding_model": EMBEDDING_MODEL,
        "index_type": "FAISS IndexFlatIP",
        "files": [
            path.name
            for path in files
        ],
    }

    with open(
        MANIFEST_FILE,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            manifest,
            file,
            ensure_ascii=False,
            indent=2
        )

    print("\n" + "=" * 60)
    print("KNOWLEDGE BASE CREATED SUCCESSFULLY")
    print("=" * 60)

    print("Documents:", len(files))
    print("Chunks:", len(metadata))
    print("Vector dimension:", dimension)

    print("\nCreated files:")

    print(INDEX_FILE)
    print(METADATA_FILE)
    print(CONFIG_FILE)
    print(MANIFEST_FILE)


if __name__ == "__main__":
    build_knowledge_base()
